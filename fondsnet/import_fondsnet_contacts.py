import csv
import hashlib
import io
import os
from dataclasses import dataclass
from datetime import UTC, datetime
from itertools import starmap
from pathlib import Path
from typing import Any, Optional, Sequence, TextIO
from urllib.parse import urlparse

import click
import openpyxl
import paramiko
import socks
import yaml
from dataclasses_json import dataclass_json
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from moneymeets.utils.collection import NonSingleValueError, get_single_value
from moneymeets.utils.s3_document import S3Document

from moneymeets_tenants.definitions import FondsnetTransactionType
from moneymeets_tenants.utils import group_by

ROOT_DIR = Path(__file__).parent.parent
FONDSNET_CONTACTS_OUTPUT_FIXTURE_PATH = Path("moneymeets_tenants/data/fixtures/fondsnet-contacts.yaml")
AB_KONFI_LIST_S3_BUCKET_AND_PATH = "it.moneymeets.net/fondsnet-konfi-lists/AB Konfi-Liste-{file_hash}.xlsx"
AB_KONFI_LIST_WEB_URL = f"https://{AB_KONFI_LIST_S3_BUCKET_AND_PATH}"
AB_KONFI_LIST_DEFAULT_DOWNLOAD_PATH = ROOT_DIR / ".tmp/AB Konfi-Liste.xlsx"
AB_KONFI_LIST_DEFAULT_DOWNLOAD_PATH.parent.mkdir(exist_ok=True)
FONDSNET_SFTP_HOST = "sftptrans.fondsnet.de"
FONDSNET_SFTP_USER = "moneymeets"
FONDSNET_SFTP_PATH = "download/AB Konfi-Liste.xlsx"

MANDANT_MONEYMEETS_USER_GROUP = "Mandant_moneymeets"


@dataclass(frozen=True)
class Row:
    ausloeser: Optional[str]
    geschaeftsart_name: Optional[str]
    geschaeftsart_id: Optional[int]
    sparte_name: Optional[str]
    sparte_id: Optional[int]
    produktgeber_name: Optional[str]
    produktgeber_id: Optional[int]
    produkt_name: Optional[str]
    produkt_id: Optional[int]
    vermittler_nummer: Optional[str]
    email: Optional[str]
    user_group: Optional[str]


@dataclass(frozen=True, order=True)
class RowContact:
    transaction_type: str
    fondsnet_company_id: int
    fondsnet_produkt_id: int
    fondsnet_geschaeftsart_id: int
    email: str
    dealer_number: str
    user_group: Optional[str]


class MultipleFondsnetContactsError(Exception):
    def __init__(
        self,
        fondsnet_company_id: int,
        fondsnet_produkt_id: int,
        transaction_type: str,
        row_contacts: Sequence[RowContact],
    ):
        self.fondsnet_company_id = fondsnet_company_id
        self.fondsnet_produkt_id = fondsnet_produkt_id
        self.transaction_type = transaction_type
        self.row_contacts = row_contacts
        super().__init__(
            f"{fondsnet_company_id=} {fondsnet_produkt_id=} {transaction_type=} {row_contacts=}",
        )


class InvalidEmail(Exception):
    def __init__(
        self,
        fondsnet_company_id: int,
        fondsnet_produkt_id: int,
        email: str,
    ):
        super().__init__(f"{fondsnet_company_id=} {fondsnet_produkt_id=} {email=}")


@dataclass_json
@dataclass(frozen=True)
class FondsnetImport:
    hash: str
    time: str


def get_konfi_list_data_from_sftp(remote_path: str) -> bytes:
    click.echo("Starting SFTP download")
    try:
        fondsnet_sftp_ssh_key = io.StringIO(os.environ["FONDSNET_SFTP_SSH_KEY"])
    except KeyError:
        fondsnet_sftp_ssh_key = (Path.home() / ".ssh/fondsnet-sftp").open()

    proxy_url = urlparse(os.environ["QUOTAGUARDSTATIC_URL"])

    def setup_socks_proxy(host: str, user: str, password: Optional[str] = None):
        socks.setdefaultproxy(
            proxy_type=socks.PROXY_TYPE_SOCKS5,
            addr=host,
            rdns=True,
            username=user,
            password=password,
        )
        paramiko.client.socket.socket = socks.socksocket

    with paramiko.SSHClient() as ssh_client:
        setup_socks_proxy(host=proxy_url.hostname, user=proxy_url.username, password=proxy_url.password)

        ssh_client.set_missing_host_key_policy(paramiko.AutoAddPolicy())

        click.echo("Establish SSH connection using SOCKS proxy")
        ssh_client.connect(
            hostname=FONDSNET_SFTP_HOST,
            username=FONDSNET_SFTP_USER,
            pkey=paramiko.RSAKey.from_private_key(fondsnet_sftp_ssh_key),
            look_for_keys=False,
            # The fondsnet server does not send server-sig-algs, thus paramiko will fail to authenticate unless we
            #   explicitly blacklist some algorithms, see also
            #   - https://stackoverflow.com/q/70565357
            #   - https://github.com/paramiko/paramiko/issues/1961
            disabled_algorithms=dict(pubkeys=["rsa-sha2-512", "rsa-sha2-256"]),
        )

        click.echo("Start SFTP client")
        sftp_client = ssh_client.open_sftp()

        with io.BytesIO() as res:
            click.echo("Reading remote file")
            sftp_client.getfo(remote_path, res)
            click.echo("Successfully retrieved data from SFTP")
            return res.getvalue()


def get_csv_from_excel(data: bytes) -> str:
    sheet = openpyxl.load_workbook(io.BytesIO(data), read_only=True)["Konfi_neu"]

    with io.StringIO() as csv_file:
        csv.writer(csv_file).writerows(sheet.iter_rows(values_only=True))
        return csv_file.getvalue()


def get_rows_from_csv(csv_file: TextIO) -> Sequence[Row]:
    def make_row(csv_row: dict) -> Row:
        def optional_str(key: str) -> Optional[str]:
            return csv_row[key].strip() or None

        def optional_int(key: str) -> Optional[int]:
            return int(csv_row[key]) if csv_row[key] else None

        return Row(
            ausloeser=optional_str("Auslöser"),
            geschaeftsart_name=optional_str("Geschäftsart"),
            geschaeftsart_id=optional_int("GA ID"),
            sparte_name=optional_str("Sparte"),
            sparte_id=optional_int("Sparte ID"),
            produktgeber_name=optional_str("Produktgeber"),
            produktgeber_id=optional_int("Produktgeber ID"),
            produkt_name=optional_str("Produkt"),
            produkt_id=optional_int("Produkt ID"),
            email=optional_str("E-Mail-Adresse"),
            vermittler_nummer=optional_str("VM-NR."),
            user_group=optional_str("User Group"),
        )

    return tuple(map(make_row, csv.DictReader(csv_file)))


def get_row_contacts_from_rows(rows: Sequence[Row]) -> Sequence[RowContact]:
    return tuple(
        RowContact(
            transaction_type=FondsnetTransactionType.CHANGE_OF_DEALER.name
            if row.ausloeser == FondsnetTransactionType.CHANGE_OF_DEALER
            else FondsnetTransactionType.ORDER.name,
            fondsnet_company_id=row.produktgeber_id,
            fondsnet_produkt_id=row.produkt_id,
            fondsnet_geschaeftsart_id=row.geschaeftsart_id,
            email=row.email or "",
            dealer_number={
                ("228-101103", 30): "759812",  # TODO: deviating HDI dealer number for order and change of dealer
            }.get((row.vermittler_nummer, row.produktgeber_id), row.vermittler_nummer or ""),
            user_group=row.user_group,
        )
        for row in rows
        if (row.user_group is None or row.user_group == MANDANT_MONEYMEETS_USER_GROUP)
        and (row.ausloeser in (FondsnetTransactionType.CHANGE_OF_DEALER, FondsnetTransactionType.ORDER))
        and (row.produktgeber_id is not None)
        and (row.produkt_id is not None)
        and not (row.email and row.email.lower().endswith("@axa-art.de"))  # TODO: MD-6046
        and not (row.email and row.email.lower().endswith("@pharmassec.de"))  # TODO: MD-6046
        and not (row.email and row.email.lower().endswith("@fondsnet.de"))
        and (row.produkt_id not in (10189, 10191))  # TODO: MD-6046
        and (
            row.vermittler_nummer not in ("58.20016.6 - keine courtagepflichtige Übertragung möglich!",)
        )  # TODO: MD-6136 duplicate dealer numbers being addressed with FONDSNET
        and (
            not (row.produkt_id == 188 and row.produktgeber_id == 8)
        )  # FONDSNET has special bKV dealer numbers for AXA, we don't support this product type
    )


def get_validated_row_contacts(row_contacts: Sequence[RowContact]) -> Sequence[RowContact]:
    def validate_contact_email(contact: RowContact):
        try:
            validate_email(contact.email)
            return contact
        except ValidationError as e:
            raise InvalidEmail(contact.fondsnet_company_id, contact.fondsnet_produkt_id, contact.email) from e

    def get_contact(group_key: tuple[int, int, str], group_row_contacts: Sequence[RowContact]) -> RowContact:
        try:
            contacts_mandant_mm = tuple(
                filter(lambda row_contact: row_contact.user_group == MANDANT_MONEYMEETS_USER_GROUP, group_row_contacts),
            )
            return get_single_value(contacts_mandant_mm or group_row_contacts)
        except NonSingleValueError:
            raise MultipleFondsnetContactsError(*group_key, group_row_contacts)

    return tuple(
        map(
            validate_contact_email,
            starmap(
                get_contact,
                group_by(
                    row_contacts,
                    key=lambda row_contact: (
                        row_contact.fondsnet_company_id,
                        row_contact.fondsnet_produkt_id,
                        row_contact.transaction_type,
                    ),
                ),
            ),
        ),
    )


def get_ab_konfi_list_url(file_hash: str) -> str:
    return AB_KONFI_LIST_WEB_URL.format(file_hash=file_hash)


def _get_contacts_fixture(data: bytes, upload: bool, current_fondsnet_import: FondsnetImport) -> str:
    contacts_data = tuple(
        {
            "fields": {
                "transaction_type": contact.transaction_type,
                "fondsnet_company_id": contact.fondsnet_company_id,
                "fondsnet_produkt_id": contact.fondsnet_produkt_id,
                "fondsnet_geschaeftsart_id": contact.fondsnet_geschaeftsart_id,
                "email": contact.email,
                "dealer_number": contact.dealer_number,
            },
            "model": "moneymeets_tenants.fondsnetcontact",
        }
        for contact in sorted(
            get_validated_row_contacts(
                sorted(set(get_row_contacts_from_rows(get_rows_from_csv(io.StringIO(get_csv_from_excel(data)))))),
            ),
        )
    )

    def get_yaml(yaml_data: Any) -> str:
        return yaml.safe_dump(yaml_data, allow_unicode=True, sort_keys=False, width=160)

    def get_fondsnet_import(contacts_content: bytes) -> FondsnetImport:
        click.echo("Comparing hashes... ", nl=False)
        new_hash = hashlib.sha256(contacts_content).hexdigest()
        if new_hash != current_fondsnet_import.hash:
            click.echo("Hash changed")
            if upload:
                _upload_file_to_s3(data, new_hash)
            else:
                click.secho("Do not commit the new hash without uploading the matching file!", fg="red")
            return FondsnetImport(new_hash, f"{datetime.now(UTC).isoformat(timespec='milliseconds')}")
        else:
            click.echo("Hash did not change")
            return current_fondsnet_import

    def get_formatted_fixture(fondsnet_import: FondsnetImport, row_contacts_data: Sequence[dict]) -> str:
        click.echo("Preparing fixtures and matching hash")
        content = get_yaml(
            (
                {
                    "fields": fondsnet_import.to_dict(),
                    "model": "moneymeets_tenants.fondsnetimport",
                },
                *row_contacts_data,
            ),
        )
        return f"# auto-generated by {Path(__file__).name}\n{content}"

    return get_formatted_fixture(
        fondsnet_import=get_fondsnet_import(get_yaml(contacts_data).encode()),
        row_contacts_data=contacts_data,
    )


def _get_current_fondsnet_import(fixtures_data: bytes) -> FondsnetImport:
    return get_single_value(
        tuple(
            FondsnetImport.from_dict(item["fields"])
            for item in yaml.safe_load(fixtures_data)
            if item["model"] == "moneymeets_tenants.fondsnetimport"
        ),
    )


def _import_data(data: bytes, upload: bool):
    click.echo("Import start")
    fixtures = _get_contacts_fixture(
        data,
        upload,
        _get_current_fondsnet_import(FONDSNET_CONTACTS_OUTPUT_FIXTURE_PATH.read_bytes()),
    )

    click.echo("Writing fixtures")
    FONDSNET_CONTACTS_OUTPUT_FIXTURE_PATH.write_text(fixtures)

    click.echo("Import end")


def _upload_file_to_s3(data: bytes, file_hash: str):
    bucket_and_path = AB_KONFI_LIST_S3_BUCKET_AND_PATH.format(file_hash=file_hash)
    click.echo(f"Uploading to {bucket_and_path}... ", nl=False)
    S3Document(bucket_and_path).upload(data, {})
    click.echo("done")


@click.group()
def main():
    pass


def upload_option():
    return click.option("--upload/--no-upload", default=False)


def path_option(exists: bool):
    return click.option(
        "--path",
        type=click.Path(exists=exists, path_type=Path),
        default=AB_KONFI_LIST_DEFAULT_DOWNLOAD_PATH,
    )


@path_option(True)
@upload_option()
@main.command("import-from-file")
def cmd_import_from_file(path: Path, upload: bool):
    _import_data(path.read_bytes(), upload)


@path_option(False)
@main.command("download-from-sftp")
def cmd_download_from_sftp(path: Path):
    path.write_bytes(get_konfi_list_data_from_sftp(FONDSNET_SFTP_PATH))


@upload_option()
@main.command("import-from-sftp")
def cmd_import_from_sftp(upload: bool):
    _import_data(get_konfi_list_data_from_sftp(FONDSNET_SFTP_PATH), upload)


if __name__ == "__main__":
    main()
